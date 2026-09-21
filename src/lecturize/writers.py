"""Output formats. Every writer starts from the same list of utterances."""

from __future__ import annotations

from pathlib import Path

from .model import Transcript
from .text import Chapter, chapters, hms

FORMATS = ("txt", "md", "srt", "docx")


def write(transcript: Transcript, out: Path, fmt: str, *, title: str | None = None) -> Path:
    if fmt not in FORMATS:
        raise ValueError(f"unknown format {fmt!r}, choose from {', '.join(FORMATS)}")
    out.parent.mkdir(parents=True, exist_ok=True)
    title = title or out.stem
    if fmt == "srt":
        write_srt(transcript, out)
    else:
        chs = chapters(transcript.utterances, transcript.language)
        {"txt": write_txt, "md": write_md, "docx": write_docx}[fmt](chs, out, title)
    return out


def write_txt(chs: list[Chapter], out: Path, title: str) -> None:
    out.write_text("\n\n".join(p for c in chs for p in c.paragraphs) + "\n", encoding="utf-8")


def write_md(chs: list[Chapter], out: Path, title: str) -> None:
    lines = [f"# {title}", ""]
    for c in chs:
        lines += [f"## {hms(c.start)}", ""]
        for p in c.paragraphs:
            lines += [p, ""]
    out.write_text("\n".join(lines), encoding="utf-8")


def write_srt(t: Transcript, out: Path) -> None:
    lines = []
    for i, u in enumerate(t.utterances, 1):
        lines += [str(i), f"{_srt_time(u.start)} --> {_srt_time(u.end)}", u.text.strip(), ""]
    out.write_text("\n".join(lines), encoding="utf-8")


def write_docx(chs: list[Chapter], out: Path, title: str) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(title, level=0)
    for c in chs:
        doc.add_heading(hms(c.start), level=2)
        for p in c.paragraphs:
            doc.add_paragraph(p)
    doc.save(str(out))


def _srt_time(seconds: float) -> str:
    total_ms = round(max(0.0, seconds) * 1000)
    h, rem = divmod(total_ms, 3_600_000)
    m, rem = divmod(rem, 60_000)
    s, ms = divmod(rem, 1000)
    return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"
